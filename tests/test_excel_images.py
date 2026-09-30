"""Offline integration tests: real xlsx files, images and drawing anchors."""
import hashlib
import importlib.util
import json
from pathlib import Path
import tempfile
import sys
import types
import unittest
from zipfile import ZipFile

import openpyxl
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
package = types.ModuleType("ff_excel_test")
package.__path__ = [str(ROOT)]
sys.modules["ff_excel_test"] = package
spec = importlib.util.spec_from_file_location("ff_excel_test.excel_images", ROOT / "excel_images.py")
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)
from ff_excel_test.xlsx_images import NS, WorkbookImages
from lxml import etree as ET


class ExcelImagesTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / "中文 表格.xlsx"
        self.image = self.root / "中文 图片.png"
        Image.new("RGBA", (300, 150), (20, 50, 90, 128)).save(self.image)
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "分镜"
        sheet.append(["镜头", "图片路径", "图片"])
        for row in range(2, 9):
            sheet.cell(row, 1, row - 1)
            sheet.cell(row, 2, self.image.name if row % 2 else str(self.image))
        sheet["D2"] = "=SUM(A2:A8)"
        sheet["A2"].number_format = "0.00"
        book.create_sheet("其他")["A1"] = "保留"
        book.save(self.source)
        book.close()
        self.original_hash = hashlib.sha256(self.source.read_bytes()).hexdigest()

    def run_node(self, **kwargs):
        options = dict(excel_path=str(self.source), image_path_column=2, start_row=2,
                       end_row=8, target_column=3, image_mode="浮动图片")
        options.update(kwargs)
        return module.FFExcelInsertImages().insert_images(**options)["result"]

    def change(self, callback):
        book = openpyxl.load_workbook(self.source)
        callback(book.active)
        book.save(self.source)
        book.close()

    def test_seven_rows_embedded_and_original_preserved(self):
        path, count, report = self.run_node()
        self.assertEqual(count, 7)
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(), self.original_hash)
        self.assertEqual(json.loads(report)["skipped"], 0)
        book = openpyxl.load_workbook(path)
        self.addCleanup(book.close)
        self.assertEqual(book["其他"]["A1"].value, "保留")
        self.assertEqual(book.active["D2"].value, "=SUM(A2:A8)")
        self.assertEqual(book.active["A2"].number_format, "0.00")
        images = book.active._images
        self.assertEqual([(i.anchor._from.col, i.anchor._from.row) for i in images], [(2, r) for r in range(1, 8)])
        for picture in images:
            self.assertAlmostEqual(picture.anchor.ext.cx / picture.anchor.ext.cy, 2, places=4)
        with ZipFile(path) as archive:
            self.assertEqual(len([n for n in archive.namelist() if n.startswith("xl/media/")]), 7)

    def test_missing_blank_invalid_keep_row_alignment(self):
        broken = self.root / "broken.png"
        broken.write_text("not an image")
        def edit(sheet):
            sheet["B3"] = "absent.png"
            sheet["B5"] = None
            sheet["B7"] = str(broken)
        self.change(edit)
        path, count, report = self.run_node()
        self.assertEqual(count, 4)
        details = json.loads(report)
        self.assertEqual([r["target"] for r in details["rows"] if r["status"] == "inserted"], ["C2", "C4", "C6", "C8"])

    def test_stop_does_not_publish_partial_output(self):
        self.change(lambda sheet: setattr(sheet["B4"], "value", "missing.png"))
        with self.assertRaisesRegex(ValueError, "B4"):
            self.run_node(on_error="stop")
        self.assertEqual(list(self.root.glob("*_with_images*")), [])

    def test_custom_destination_sheet_and_offset(self):
        path, count, report = self.run_node(sheet_name="分镜", target_start_row=10,
                                           target_column=5, output_path="新文件夹/result.xlsx")
        self.assertEqual(Path(path).parent, self.root / "新文件夹")
        self.assertEqual(json.loads(report)["rows"][-1]["target"], "E16")
        book = openpyxl.load_workbook(path)
        self.addCleanup(book.close)
        self.assertEqual(book.active._images[0].anchor._from.row, 9)

    def test_unique_output_never_overwrites(self):
        first = self.run_node()[0]
        digest = hashlib.sha256(Path(first).read_bytes()).hexdigest()
        second = self.run_node()[0]
        self.assertNotEqual(first, second)
        self.assertEqual(hashlib.sha256(Path(first).read_bytes()).hexdigest(), digest)
        with self.assertRaisesRegex(ValueError, "原表"):
            self.run_node(output_path=str(self.source))

    def test_rerun_replaces_target_images(self):
        first = self.run_node()[0]
        second = self.run_node(excel_path=first)[0]
        book = openpyxl.load_workbook(second)
        self.addCleanup(book.close)
        self.assertEqual(len(book.active._images), 7)

    def test_merged_target_rejected(self):
        self.change(lambda sheet: sheet.merge_cells("C2:C3"))
        with self.assertRaisesRegex(ValueError, "合并"):
            self.run_node()

    def test_wps_embedding_and_existing_images_preserved(self):
        first, count, _ = self.run_node(image_mode="嵌入单元格（WPS）")
        self.assertEqual(count, 7)
        with ZipFile(first) as archive:
            root = ET.fromstring(archive.read("xl/cellimages.xml"))
            self.assertEqual(len(root), 7)
            ids = [n.get("name") for n in root.findall(".//xdr:cNvPr", NS)]
            sheet = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
            self.assertIsNone(sheet.find("s:drawing", NS))
            formulas = [sheet.find(f"s:sheetData/s:row/s:c[@r='C{r}']/s:f", NS).text for r in range(2, 9)]
            for identifier, formula in zip(ids, formulas):
                self.assertEqual(formula, f'_xlfn.DISPIMG("{identifier}",1)')
            types = ET.fromstring(archive.read("[Content_Types].xml"))
            cell_type = next(n for n in types if n.get("PartName") == "/xl/cellimages.xml")
            self.assertEqual(cell_type.get("ContentType"), "application/vnd.wps-officedocument.cellimage+xml")
            original_media = {n: archive.read(n) for n in archive.namelist() if n.startswith("xl/media/")}
        second = self.run_node(excel_path=first, image_mode="嵌入单元格（WPS）", target_column=5)[0]
        with ZipFile(second) as archive:
            self.assertEqual(len(ET.fromstring(archive.read("xl/cellimages.xml"))), 14)
            for name, content in original_media.items():
                self.assertEqual(archive.read(name), content)

    def test_grid_dimensions_unchanged_in_both_modes(self):
        def edit(sheet):
            sheet.column_dimensions["C"].width = 9.5
            for r in range(2, 9):
                sheet.row_dimensions[r].height = 18 + r * 3
        self.change(edit)
        with ZipFile(self.source) as archive:
            before = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        for mode in ("嵌入单元格（WPS）", "浮动图片"):
            with self.subTest(mode=mode):
                result = self.run_node(image_mode=mode)[0]
                with ZipFile(result) as archive:
                    after = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
                for element in ("cols", "sheetFormatPr"):
                    self.assertEqual(ET.tostring(before.find("s:" + element, NS)), ET.tostring(after.find("s:" + element, NS)))
                self.assertEqual([dict(r.attrib) for r in before.findall("s:sheetData/s:row", NS)], [dict(r.attrib) for r in after.findall("s:sheetData/s:row", NS)])

    def test_floating_fits_original_cells(self):
        self.change(lambda sheet: setattr(sheet.row_dimensions[2], "height", 21))
        path = self.run_node()[0]
        book = openpyxl.load_workbook(path)
        self.addCleanup(book.close)
        for picture in book.active._images:
            anchor = picture.anchor
            row = anchor._from.row + 1
            max_height = (book.active.row_dimensions[row].height or 15) * 96 / 72
            self.assertLessEqual((anchor._from.rowOff + anchor.ext.cy) / 9525, max_height - 1.9)
            self.assertLessEqual((anchor._from.colOff + anchor.ext.cx) / 9525, 64 - 1.9)

    def test_mode_switch_replaces_target_picture(self):
        embedded = self.run_node(image_mode="嵌入单元格（WPS）")[0]
        floating = self.run_node(excel_path=embedded)[0]
        book = openpyxl.load_workbook(floating)
        self.addCleanup(book.close)
        self.assertEqual(len(book.active._images), 7)
        self.assertIsNone(book.active["C2"].value)
        final = self.run_node(excel_path=floating, image_mode="嵌入单元格（WPS）")[0]
        with ZipFile(final) as archive:
            for name in archive.namelist():
                if name.startswith("xl/drawings/") and name.endswith(".xml"):
                    self.assertEqual(len(ET.fromstring(archive.read(name)).findall(".//xdr:pic", NS)), 0)

    def test_untouched_parts_preserved_byte_for_byte(self):
        with ZipFile(self.source, "a") as archive:
            archive.writestr("customXml/item1.xml", '<keep xmlns="urn:ff-test"/>')
        output = self.run_node(image_mode="嵌入单元格（WPS）")[0]
        with ZipFile(self.source) as before, ZipFile(output) as after:
            for name in ("xl/styles.xml", "xl/worksheets/sheet2.xml", "xl/workbook.xml", "customXml/item1.xml"):
                self.assertEqual(before.read(name), after.read(name))

    def test_hidden_rows_skipped_without_unhiding(self):
        self.change(lambda sheet: setattr(sheet.row_dimensions[3], "hidden", True))
        path, count, report = self.run_node()
        self.assertEqual(count, 6)
        self.assertEqual(json.loads(report)["rows"][1]["status"], "skipped")
        book = openpyxl.load_workbook(path)
        self.addCleanup(book.close)
        self.assertTrue(book.active.row_dimensions[3].hidden)

    def test_overlapping_path_and_destination_snapshot(self):
        _, count, _ = self.run_node(target_column=2, target_start_row=3, image_mode="嵌入单元格（WPS）")
        self.assertEqual(count, 7)

    def test_parameters_rejected(self):
        for args in ({"end_row": 1}, {"image_path_column": 0}, {"sheet_name": "不存在"},
                     {"target_start_row": 1048576}, {"padding": -1}, {"image_mode": "unknown"},
                     {"on_error": "unknown"}, {"output_path": "out.xls"}):
            with self.subTest(args=args), self.assertRaises(ValueError):
                self.run_node(**args)

    def test_formula_path_is_not_executed(self):
        self.change(lambda sheet: setattr(sheet["B2"], "value", '=HYPERLINK("x")'))
        _, count, report = self.run_node()
        self.assertEqual(count, 6)
        self.assertIn("公式", json.loads(report)["rows"][0]["reason"])

    def test_all_invalid_produces_no_output(self):
        self.change(lambda sheet: [setattr(sheet.cell(r, 2), "value", None) for r in range(2, 9)])
        with self.assertRaisesRegex(ValueError, "没有可插入"):
            self.run_node()
        self.assertEqual(list(self.root.glob("*_with_images*")), [])


if __name__ == "__main__":
    unittest.main()
